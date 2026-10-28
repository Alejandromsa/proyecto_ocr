"""Integración sobre los PDF ficticios del mockup y copias alteradas generadas al vuelo (verdad conocida)."""
from __future__ import annotations

import io
import zipfile

import pikepdf
import pytest
from PIL import Image

from stracon_ocr.contrato import Documento, esquema_json
from stracon_ocr.pipeline import Procesador

from .conftest import HOY, mockup

LIMPIO = "16884102-7/6-certificado-antecedentes.pdf"


def _ids(d: Documento, minimo: str = "media") -> set[str]:
    orden = {"info": 0, "baja": 1, "media": 2, "alta": 3}
    return {s.id for s in d.senales if orden[s.severidad] >= orden[minimo]}


def test_documento_limpio_sin_senales_y_campos(proc_sin_bd: Procesador):
    d = proc_sin_bd.procesar(mockup(LIMPIO), expediente="16884102-7", guardar=False, hoy=HOY)
    assert d.clasificacion.tipo == "certificado_antecedentes"
    assert d.archivo.origen == "pdf_nativo"
    assert _ids(d) == set(), [s.detalle for s in d.senales]
    assert d.campos["rut"].valor_normalizado == "16884102-7"
    assert d.campos["nombre"].valor_normalizado == "BENJAMIN SILVA ARAYA"
    assert d.campos["resultado"].valor == "Sin antecedentes"
    assert all(v.resultado != "falla" for v in d.validaciones)
    # el JSON cumple el contrato (ida y vuelta) y el esquema se puede exportar
    assert Documento.model_validate_json(d.model_dump_json()).id == d.id
    assert "documento" in esquema_json()


@pytest.mark.parametrize("archivo, senal, campo", [
    ("12794536-5/2-licencia-conducir.pdf", "pdf.fuentes_mezcladas_campo", "folio"),
    ("15419663-3/4-certificado-ceim.pdf", "pdf.texto_superpuesto", "fecha_emision"),
    ("18997304-7/6-certificado-antecedentes.pdf", "pdf.campo_rasterizado", "codigo_verificacion"),
    ("18452307-8/4-certificado-ceim.pdf", "pdf.imagen_baja_nitidez", None),
])
def test_alteraciones_del_mockup(proc_sin_bd: Procesador, archivo, senal, campo):
    d = proc_sin_bd.procesar(mockup(archivo), guardar=False, hoy=HOY)
    hallazgos = [s for s in d.senales if s.id == senal]
    assert hallazgos, [s.id for s in d.senales]
    if campo:
        assert any(s.campo == campo for s in hallazgos)
    assert all(s.bbox and s.pagina for s in hallazgos), "cada hallazgo debe ubicarse en la página"


def test_superposicion_muestra_ambas_capas(proc_sin_bd: Procesador):
    d = proc_sin_bd.procesar(mockup("15419663-3/4-certificado-ceim.pdf"), guardar=False, hoy=HOY)
    s = next(s for s in d.senales if s.id == "pdf.texto_superpuesto")
    assert s.evidencia["visible"].replace(" ", "") == "20sep2026"
    assert s.evidencia["debajo"].replace(" ", "") == "20sep2024"
    assert d.campos["fecha_emision"].valor_normalizado == "2026-09-20"


def test_revision_incremental_y_productor_editor(proc_sin_bd: Procesador, tmp_path):
    pypdf = pytest.importorskip("pypdf")
    original = mockup(LIMPIO)
    buf = io.BytesIO()
    w = pypdf.PdfWriter(str(original), incremental=True)
    w.add_metadata({"/Producer": "iLovePDF", "/ModDate": "D:20261020180000+00'00'"})
    w.write(buf)
    ruta = tmp_path / "alterado.pdf"
    ruta.write_bytes(buf.getvalue())
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    ids = _ids(d)
    assert "pdf.revisiones_incrementales" in ids
    assert "meta.productor_no_esperado" in ids
    assert d.archivo.metadatos["revisiones"]["revisiones_extra"] >= 1


def test_texto_tapado_por_rectangulo_blanco(proc_sin_bd: Procesador, tmp_path):
    """Un rectángulo blanco dibujado encima del RUT: el texto sigue en el PDF pero no se ve."""
    pdf = pikepdf.open(mockup(LIMPIO))
    pagina = pdf.pages[0]
    alto = float(pagina.mediabox[3])
    # RUT en coordenadas pdfplumber (arriba-izquierda) ≈ x 315-400, top 234-248 → PDF (abajo-izquierda)
    y0, y1 = alto - 250, alto - 232
    # el contenido original deja una transformación activa: se encierra en q … Q antes de dibujar encima
    pagina.contents_add(pikepdf.Stream(pdf, b"q\n"), prepend=True)
    pagina.contents_add(pikepdf.Stream(pdf, f"\nQ q 1 1 1 rg 312 {y0} 95 {y1 - y0} re f Q".encode()), prepend=False)
    ruta = tmp_path / "tapado.pdf"
    pdf.save(ruta)
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    tapas = [s for s in d.senales if s.id == "pdf.objeto_tapa_texto"]
    assert tapas and tapas[0].severidad == "alta", [s.id for s in d.senales]


def test_foto_retocada_y_ocr_de_imagen(proc_sin_bd: Procesador, tmp_path):
    """Página renderizada a JPG con EXIF de Photoshop: entra por el camino OCR y extrae los campos."""
    import pypdfium2 as pdfium

    img = pdfium.PdfDocument(str(mockup(LIMPIO)))[0].render(scale=200 / 72).to_pil().convert("RGB")
    exif = Image.Exif()
    exif[305] = "Adobe Photoshop 25.0 (Windows)"
    ruta = tmp_path / "foto.jpg"
    img.save(ruta, "JPEG", quality=88, exif=exif)
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    assert d.archivo.origen == "imagen" and d.paginas[0].fuente_texto == "ocr"
    assert d.clasificacion.tipo == "certificado_antecedentes"
    assert d.campos["rut"].valor_normalizado == "16884102-7"
    assert "img.exif_editor" in _ids(d)


def test_pdf_escaneado_de_documento_nativo(proc_sin_bd: Procesador, tmp_path):
    import pypdfium2 as pdfium

    img = pdfium.PdfDocument(str(mockup(LIMPIO)))[0].render(scale=200 / 72).to_pil().convert("RGB")
    ruta = tmp_path / "escaneado.pdf"
    img.save(ruta, "PDF", resolution=200)
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    assert d.archivo.origen == "pdf_escaneado"
    # escanear o imprimir a PDF un certificado electrónico es habitual entre postulantes (calibración oct 2026: 10 de 133
    # documentos reales): la señal queda en baja y lo que se pide es el original, con la observación
    solo = [s for s in d.senales if s.id == "pdf.solo_imagen"]
    assert solo and solo[0].severidad == "baja"
    assert any(o.id == "contenido.imagen_de_electronico" and o.accion == "solicitar_original" for o in d.observaciones)
    assert d.campos["folio"].valor_normalizado == "700-168-4102"


def test_reintento_a_mayor_dpi_cuando_faltan_campos_obligatorios(proc_sin_bd: Procesador, tmp_path, monkeypatch):
    """A baja resolución el OCR puede perder una línea (p. ej. el folio): el PDF escaneado se vuelve a leer a más DPI."""
    import pypdfium2 as pdfium
    from stracon_ocr import pipeline

    img = pdfium.PdfDocument(str(mockup(LIMPIO)))[0].render(scale=200 / 72).to_pil().convert("RGB")
    ruta = tmp_path / "escaneado.pdf"
    img.save(ruta, "PDF", resolution=200)
    original = pipeline.extraccion.extraer

    def extraer_ciego(doc, tipo):  # simula un OCR que a baja resolución no lee el folio
        campos = original(doc, tipo)
        if doc.paginas and doc.paginas[0].imagen.width < 2000:
            campos.pop("folio", None)
        return campos

    monkeypatch.setattr(pipeline.extraccion, "extraer", extraer_ciego)
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    assert "folio" in d.campos and "reintento_dpi" in d.tiempos_ms
    assert not any(o.id == "contenido.campos_faltantes" for o in d.observaciones)
    # sin reintento (SOCR_DPI_REINTENTO=0) se conserva la lectura original y se pide el dato faltante
    proc_sin_bd.cfg.dpi_reintento = 0
    try:
        d0 = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    finally:
        proc_sin_bd.cfg.dpi_reintento = 400
    assert "folio" not in d0.campos and "reintento_dpi" not in d0.tiempos_ms
    assert any(o.id == "contenido.campos_faltantes" for o in d0.observaciones)


def test_word_con_cambios_rastreados(proc_sin_bd: Procesador, tmp_path):
    from docx import Document as Docx

    doc = Docx()
    doc.add_heading("Certificado de Antecedentes", 0)
    for linea in ("Servicio de Registro Civil e Identificación", "Nombre: Benjamín Silva Araya", "RUT: 16.884.102-7",
                  "Folio: 700-168-4102", "Fecha de emisión: 5 oct 2026", "Resultado: Sin antecedentes"):
        doc.add_paragraph(linea)
    buf = io.BytesIO()
    doc.save(buf)
    # inyecta una inserción rastreada
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w") as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename == "word/document.xml":
                datos = datos.replace(b"</w:body>", b'<w:p><w:ins w:id="1" w:author="X" w:date="2026-10-20T10:00:00Z"><w:r><w:t>agregado</w:t></w:r></w:ins></w:p></w:body>')
            zout.writestr(item, datos)
    ruta = tmp_path / "antecedentes.docx"
    ruta.write_bytes(salida.getvalue())
    d = proc_sin_bd.procesar(ruta, guardar=False, hoy=HOY)
    assert d.archivo.origen == "docx"
    assert d.clasificacion.tipo == "certificado_antecedentes"
    assert d.campos["rut"].valor_normalizado == "16884102-7"
    assert d.campos["fecha_emision"].valor_normalizado == "2026-10-05"
    assert {"docx.certificado_en_word", "docx.cambios_rastreados"} <= _ids(d)


def test_expediente_contraste_y_duplicados(cfg_tmp):
    proc = Procesador(cfg_tmp)
    rut = "12794536-5"
    for nombre in ("1-cedula-identidad.pdf", "2-licencia-conducir.pdf", "3-hoja-vida-conductor.pdf"):
        proc.procesar(mockup(f"{rut}/{nombre}"), expediente=rut, hoy=HOY)
    exp = proc.contrastar(rut)
    assert exp.identidad.rut == rut
    assert set(exp.faltantes) == {"certificado_estudios_medios", "certificado_estudios_profesionales", "certificado_antecedentes"}
    assert not [x for x in exp.discrepancias if x.campo in {"rut", "nombre"}]
    assert any(f.documento == "hoja_vida_conductor" and f.estado == "pendiente_revisor" for f in exp.fuentes)
    # el mismo archivo entregado por otro postulante
    d = proc.procesar(mockup(f"{rut}/1-cedula-identidad.pdf"), expediente="99999999-9", hoy=HOY)
    assert "dup.archivo_identico" in _ids(d)
